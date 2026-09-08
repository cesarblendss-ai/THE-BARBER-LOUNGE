#!/usr/bin/env python3
"""Apply known off-Amex buildout adjustments and refresh totals."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
CLASSIFIED = ROOT / "amex-buildout-classified.xlsx"
PARTNERSHIP = ROOT / "partnership-split-model.xlsx"

# Manual known items Cesar confirmed (not on / not fully on Amex exports)
# Adjust overlap_note if Cesar says $11,600 INCLUDES the Amex FNR charges.
ADJUSTMENTS = [
    {
        "date": "2026-02/03 (approx)",
        "amount": 11600.00,
        "bucket": "Buildout",
        "buildout": "Yes",
        "category": "Construction / contractor / permits",
        "member": "Alexis (cash/transfer to pay contractor)",
        "desc": "Cash/transfer so Alexis can pay contractor",
        "note": (
            "Cesar confirmed $11,600 for Alexis → contractor. "
            "Assumed IN ADDITION to Amex FNR charges ($2,000 Cesar + $2,350 Alexis = $4,350). "
            "Change if $11,600 already includes those Amex hits."
        ),
        "source": "manual-cesar",
        "ref": "MANUAL-ALEXIS-CONTRACTOR-11600",
    }
]

thin = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
header_fill = PatternFill("solid", fgColor="1A1A1A")
header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
warn_fill = PatternFill("solid", fgColor="F8D7DA")
ok_fill = PatternFill("solid", fgColor="D1E7DD")
calc_fill = PatternFill("solid", fgColor="E8F4EA")
input_fill = PatternFill("solid", fgColor="FFF3CD")
money = '"$"#,##0.00'


def style_header(ws, row, start, end):
    for col in range(start, end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin


def main():
    wb = load_workbook(CLASSIFIED)

    # ---- ManualAdjustments sheet ----
    if "ManualAdjustments" in wb.sheetnames:
        del wb["ManualAdjustments"]
    ma = wb.create_sheet("ManualAdjustments", 1)
    headers = [
        "Date",
        "Amount",
        "Bucket",
        "Buildout?",
        "Category",
        "Who / how paid",
        "Description",
        "Note",
        "Source",
        "Reference",
    ]
    for i, h in enumerate(headers, 1):
        ma.cell(row=1, column=i, value=h)
    style_header(ma, 1, 1, len(headers))

    for i, adj in enumerate(ADJUSTMENTS, 2):
        vals = [
            adj["date"],
            adj["amount"],
            adj["bucket"],
            adj["buildout"],
            adj["category"],
            adj["member"],
            adj["desc"],
            adj["note"],
            adj["source"],
            adj["ref"],
        ]
        for ci, v in enumerate(vals, 1):
            cell = ma.cell(row=i, column=ci, value=v)
            cell.border = thin
            cell.fill = input_fill
            if ci == 2:
                cell.number_format = money

    # Append to BuildoutOnly if not already present
    bo = wb["BuildoutOnly"]
    existing_refs = set()
    for r in range(2, bo.max_row + 1):
        existing_refs.add(str(bo.cell(row=r, column=11).value or ""))

    next_row = bo.max_row + 1
    for adj in ADJUSTMENTS:
        if adj["ref"] in existing_refs:
            continue
        vals = [
            adj["date"],
            adj["amount"],
            adj["bucket"],
            adj["buildout"],
            adj["category"],
            adj["member"],
            adj["desc"],
            adj["note"],
            "",
            adj["source"],
            adj["ref"],
        ]
        for ci, v in enumerate(vals, 1):
            cell = bo.cell(row=next_row, column=ci, value=v)
            cell.border = thin
            cell.fill = warn_fill
            if ci == 2:
                cell.number_format = money
        next_row += 1

    # Also append to AllLines
    al = wb["AllLines"]
    existing_refs_al = set()
    for r in range(2, al.max_row + 1):
        existing_refs_al.add(str(al.cell(row=r, column=11).value or ""))
    next_row = al.max_row + 1
    for adj in ADJUSTMENTS:
        if adj["ref"] in existing_refs_al:
            continue
        vals = [
            adj["date"],
            adj["amount"],
            adj["bucket"],
            adj["buildout"],
            adj["category"],
            adj["member"],
            adj["desc"],
            adj["note"],
            "",
            adj["source"],
            adj["ref"],
        ]
        for ci, v in enumerate(vals, 1):
            cell = al.cell(row=next_row, column=ci, value=v)
            cell.border = thin
            cell.fill = warn_fill
            if ci == 2:
                cell.number_format = money
        next_row += 1

    # Recompute headline from BuildoutOnly Yes rows + Manual
    yes_total = 0.0
    generous = 0.0
    by_cat = {}
    for r in range(2, bo.max_row + 1):
        amt = bo.cell(row=r, column=2).value or 0
        flag = bo.cell(row=r, column=4).value
        cat = bo.cell(row=r, column=5).value or "Other"
        if flag == "Yes":
            yes_total += float(amt)
            by_cat[cat] = by_cat.get(cat, 0) + float(amt)
            generous += float(amt)
        elif flag == "Review":
            generous += float(amt)

    # Read ops/personal from summary sheet (unchanged Amex-only)
    ws = wb["BuildoutSummary"]
    ops = ws["B8"].value or 0
    personal = ws["B9"].value or 0
    review = ws["B10"].value or 0

    ws["B6"] = round(yes_total, 2)
    ws["B6"].number_format = money
    ws["B6"].fill = warn_fill
    ws["C6"] = "Amex Yes + manual contractor $11,600 (assumed additive to Amex FNR)"

    ws["B7"] = round(generous, 2)
    ws["B7"].number_format = money
    ws["B7"].fill = warn_fill

    # Update contractor category line if present, else leave category block stale —
    # rewrite category block from by_cat
    ws["A12"] = "BUILDOUT BY CATEGORY (Yes + manual)"
    # clear old category rows 14-25
    for r in range(14, 26):
        ws.cell(row=r, column=1).value = None
        ws.cell(row=r, column=2).value = None
    for i, (cat, val) in enumerate(sorted(by_cat.items(), key=lambda x: -abs(x[1])), 14):
        ws.cell(row=i, column=1, value=cat).border = thin
        cell = ws.cell(row=i, column=2, value=round(val, 2))
        cell.number_format = money
        cell.fill = calc_fill
        cell.border = thin

    ws["A28"] = "Manual add"
    ws["A28"].font = Font(name="Calibri", bold=True, size=12)
    ws["A29"] = (
        "$11,600 to Alexis for contractor. Amex already shows FNR $4,350 "
        "($2,000 Cesar + $2,350 Alexis). Current model treats $11,600 as EXTRA. "
        "If the $11,600 was the total contractor budget and already covers those Amex charges, "
        "true added cash is $11,600 − $4,350 = $7,250 — tell me and I'll switch."
    )
    ws["A29"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A29:E29")
    ws.row_dimensions[29].height = 70

    wb.save(CLASSIFIED)

    # Patch partnership Buildout headline if sheet exists
    if PARTNERSHIP.exists():
        pwb = load_workbook(PARTNERSHIP)
        if "Buildout" in pwb.sheetnames:
            bd = pwb["Buildout"]
            bd["B6"] = round(yes_total, 2)
            bd["B6"].number_format = money
            bd["B6"].fill = warn_fill
            bd["A6"] = "BUILDOUT confident (Amex Yes + $11,600 contractor)"
            bd["B7"] = round(generous, 2)
            bd["B7"].number_format = money
            # append manual line if missing
            found = False
            for r in range(14, bd.max_row + 1):
                if bd.cell(row=r, column=8).value and "MANUAL-ALEXIS-CONTRACTOR" in str(
                    bd.cell(row=r, column=8).value
                ):
                    found = True
                    break
                # also check notes col
                if "Alexis can pay contractor" in str(bd.cell(row=r, column=3).value or ""):
                    found = True
                    break
            if not found:
                r = bd.max_row + 1
                vals = [
                    "2026-02/03",
                    "Alexis (cash/transfer)",
                    "Cash/transfer so Alexis can pay contractor",
                    11600.0,
                    "Construction / contractor / permits",
                    "Yes",
                    "Cesar→Alexis",
                    "Manual add — confirm not double-counting Amex FNR $4,350",
                ]
                for ci, v in enumerate(vals, 1):
                    cell = bd.cell(row=r, column=ci, value=v)
                    cell.border = thin
                    cell.fill = warn_fill
                    if ci == 4:
                        cell.number_format = money
            pwb.save(PARTNERSHIP)

    print(f"Buildout Yes total now: ${yes_total:,.2f}")
    print(f"Contractor category: ${by_cat.get('Construction / contractor / permits', 0):,.2f}")
    print(f"Generous: ${generous:,.2f}")
    print("Assumed $11,600 ADDITIVE to Amex FNR $4,350")


if __name__ == "__main__":
    main()
